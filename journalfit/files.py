"""Journal-ready files made from the author's own manuscript, as a careful author would:

* title_page.docx: title, authors, affiliations, corresponding author, ORCIDs and the
  declarations (always made; most portals have a Title Page file type);
* anonymised_manuscript.docx for double-blind journals: the author block, acknowledgements,
  funding, contributions and file properties removed, and names/emails/institutions blanked;
* figures/Figure_N.ext: the pictures inside the Word file, in order, when the journal wants
  separate figure files;
* highlights.docx/.txt when the journal requires highlights (written by the AI from the
  abstract, for the author to check).

The original manuscript is never changed.
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from .submit import people

BLIND_SECTIONS = re.compile(r"^\s*(\d+(\.\d+)*\.?\s*)?(acknowledge?ments?|funding(\s+(information|statement|sources?))?|"
                            r"financial\s+support|authors?'?\s+contributions?|credit\s+(authorship\s+)?(contribution\s+)?"
                            r"statement|author\s+information|ethics\s+declarations?|declarations?)\s*:?\s*$", re.I)
SECTION_HEAD = re.compile(r"^\s*(\d+(\.\d+)*\.?\s+)?[A-Z][^.!?]{0,80}$")
ABSTRACT = re.compile(r"^\s*(abstract|summary)\b", re.I)
REFERENCES = re.compile(r"^\s*(\d+\.?\s*)?(references|bibliography|literature\s+cited|works\s+cited)\s*:?\s*$", re.I)


def title_page(path: Path, title: str, authors: list[dict], declarations: dict, word_count: int = 0) -> Path:
    import docx

    d = docx.Document()
    d.add_heading(title, level=1)
    affs: list[str] = []
    names = []
    for a in authors:
        aff = ", ".join(x for x in (a.get("department"), a.get("institution"), a.get("city"), a.get("state"),
                                   a.get("country")) if x)
        if aff and aff not in affs:
            affs.append(aff)
        mark = str(affs.index(aff) + 1) if aff else ""
        names.append(people.full_name(a) + (f"^{mark}" if mark else "") + ("*" if a.get("corresponding") else ""))
    p = d.add_paragraph()
    for i, n in enumerate(names):
        base, _, sup = n.partition("^")
        p.add_run(base)
        if sup:
            p.add_run(sup).font.superscript = True
        elif n.endswith("*"):
            pass
        if i < len(names) - 1:
            p.add_run(", ")
    for i, aff in enumerate(affs, 1):
        q = d.add_paragraph()
        q.add_run(str(i)).font.superscript = True
        q.add_run(" " + aff)
    corr = next((a for a in authors if a.get("corresponding")), None)
    if corr:
        d.add_paragraph("*Corresponding author: " + people.full_name(corr)
                        + (f", {corr['email']}" if corr.get("email") else ""))
    orcids = [f"{people.full_name(a)}: https://orcid.org/{a['orcid']}" for a in authors if a.get("orcid")]
    if orcids:
        d.add_paragraph("ORCID iDs: " + "; ".join(orcids))
    if word_count:
        d.add_paragraph(f"Word count: {word_count}")
    for key, head in (("acknowledgements", "Acknowledgements"), ("funding", "Funding"),
                      ("competing_interests", "Declaration of competing interests"),
                      ("author_contributions", "Author contributions"), ("data_availability", "Data availability"),
                      ("ethics", "Ethics approval"), ("ai_disclosure", "Use of generative AI")):
        if declarations.get(key):
            d.add_heading(head, level=2)
            d.add_paragraph(declarations[key])
    d.save(str(path))
    return path


def anonymise(src: Path, dest: Path, authors: list[dict]) -> list[str]:
    """Copy the Word manuscript without anything that identifies the authors. Returns what was removed."""
    import docx

    shutil.copy2(src, dest)
    d = docx.Document(str(dest))
    removed: list[str] = []
    paras = d.paragraphs
    # 1. The author block: everything between the title and the abstract.
    first = next((i for i, p in enumerate(paras[:15]) if len(p.text.strip()) >= 15), None)
    abstract = next((i for i, p in enumerate(paras[:60]) if ABSTRACT.match(p.text)), None)
    if first is not None and abstract is not None and abstract > first + 1:
        for p in paras[first + 1:abstract]:
            if p.text.strip():
                removed.append("author block: " + p.text.strip()[:80])
            _drop(p)
    # 2. Sections that name the authors or their funders.
    paras = d.paragraphs
    dropping = False
    for p in paras:
        t = p.text.strip()
        if REFERENCES.match(t):
            break
        if BLIND_SECTIONS.match(t) or re.match(r"^\s*(acknowledge?ments?|funding)\s*:", t, re.I):
            dropping = True
            removed.append("section: " + t[:60])
            _drop(p)
            continue
        if dropping:
            if t and SECTION_HEAD.match(t) and len(t) < 80 and not t.endswith(".") and not BLIND_SECTIONS.match(t):
                dropping = False
            else:
                _drop(p)
                continue
    # 3. Names, emails and institutions left in the text (references are kept as they are).
    needles = set()
    for a in authors:
        for v in (people.full_name(a), a.get("email"), a.get("orcid")):
            if v and len(v) > 4:
                needles.add(v)
        if a.get("institution") and len(a["institution"]) > 8:
            needles.add(a["institution"])
    blanked = 0
    for p in d.paragraphs:
        if REFERENCES.match(p.text.strip()):
            break
        for run in p.runs:
            for n in sorted(needles, key=len, reverse=True):
                if n.lower() in run.text.lower():
                    run.text = re.sub(re.escape(n), "[blinded]", run.text, flags=re.I)
                    blanked += 1
    if blanked:
        removed.append(f"{blanked} mention(s) of author names, emails or institutions replaced by [blinded]")
    # 4. File properties (Word stores the author's name in the file).
    cp = d.core_properties
    for attr in ("author", "last_modified_by", "comments", "title", "subject", "keywords", "category"):
        if getattr(cp, attr, None):
            setattr(cp, attr, "")
    removed.append("file properties (author, last modified by)")
    for section in d.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if any(n.lower() in p.text.lower() for n in needles):
                    _drop(p)
                    removed.append("header/footer line with an author's name")
    d.save(str(dest))
    return removed


def _drop(p) -> None:
    el = p._element
    el.getparent().remove(el)


def extract_figures(src: Path, folder: Path) -> list[Path]:
    """Pictures inside the Word file, saved as Figure_1, Figure_2, ... in the order they appear."""
    import docx

    d = docx.Document(str(src))
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    seen, out = [], []
    for el in d.element.body.iter():
        if el.tag.endswith("}blip") and el.get(ns) and el.get(ns) not in seen:
            seen.append(el.get(ns))
    for i, rid in enumerate(seen, 1):
        part = d.part.related_parts.get(rid)
        if part is None:
            continue
        ext = Path(part.partname).suffix or ".png"
        folder.mkdir(parents=True, exist_ok=True)
        dest = folder / f"Figure_{i}{ext}"
        dest.write_bytes(part.blob)
        out.append(dest)
    return out


def figure_sizes(paths: list[Path]) -> list[str]:
    """'Figure_1.png: 2400x1600 px' — so the author can see whether figures are print quality."""
    out = []
    for p in paths:
        size = _image_size(p.read_bytes())
        out.append(f"{p.name}: {size[0]}x{size[1]} px" if size else p.name)
    return out


def _image_size(b: bytes):
    import struct

    if b[:8] == b"\x89PNG\r\n\x1a\n":
        return struct.unpack(">II", b[16:24])
    if b[:2] == b"\xff\xd8":
        i = 2
        while i < len(b) - 9:
            if b[i] != 0xFF:
                i += 1
                continue
            marker = b[i + 1]
            if marker in (0xC0, 0xC1, 0xC2):
                h, w = struct.unpack(">HH", b[i + 5:i + 9])
                return w, h
            i += 2 + struct.unpack(">H", b[i + 2:i + 4])[0]
    return None


def highlights_files(folder: Path, bullets: list[str]) -> Path:
    import docx

    bullets = [b.strip().lstrip("•-* ").strip() for b in bullets if b.strip()]
    (folder / "highlights.txt").write_text("\n".join(bullets), encoding="utf-8")
    d = docx.Document()
    d.add_heading("Highlights", level=2)
    for b in bullets:
        d.add_paragraph(b, style="List Bullet")
    d.save(str(folder / "highlights.docx"))
    return folder / "highlights.docx"
