"""Read a manuscript (DOCX, PDF, LaTeX, Markdown or plain text) into a Manuscript.

Extraction is heuristic and deliberately conservative; the LLM profiler
(llm.profile_manuscript) can fill gaps such as a missing title or keywords.
"""

from __future__ import annotations

import re
from pathlib import Path

from .models import Manuscript

DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>,;]+)", re.I)
HEADING_REF_RE = re.compile(
    r"^\s*(?:\d+\.?\s*)?(references|bibliography|literature cited|works cited)\s*:?\s*$", re.I
)
HEADING_ABSTRACT_RE = re.compile(r"^\s*(abstract|summary)\s*[:.\-–—]?\s*(.*)$", re.I)
KEYWORDS_RE = re.compile(r"^\s*(key\s*words?|index terms)\s*[:.\-–—]\s*(.+)$", re.I)
SECTION_AFTER_ABSTRACT_RE = re.compile(
    r"^\s*(?:\d+\.?\s*)?(introduction|background|keywords?|key words|index terms|1\s)\b", re.I
)
FIGURE_RE = re.compile(r"\b(?:fig(?:ure)?\.?)\s*(\d+)", re.I)
TABLE_RE = re.compile(r"\btable\s*(\d+)", re.I)


def read_text(path: Path) -> list[str]:
    """Return the document as a list of paragraphs/lines."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        paras = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paras.append(" | ".join(c.text for c in row.cells))
        return paras
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return text.splitlines()
    if suffix in (".tex", ".latex"):
        return _latex_to_lines(path.read_text(encoding="utf-8", errors="replace"))
    return path.read_text(encoding="utf-8", errors="replace").splitlines()


def _latex_to_lines(src: str) -> list[str]:
    src = re.sub(r"(?<!\\)%.*", "", src)
    lines: list[str] = []
    m = re.search(r"\\title(?:\[[^\]]*\])?\{(.+?)\}\s*$", src, re.S | re.M)
    if m:
        lines.append(_strip_latex(m.group(1)))
    m = re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", src, re.S)
    if m:
        lines.append("Abstract")
        lines.append(_strip_latex(m.group(1)))
    m = re.search(r"\\(?:keywords|begin\{keyword\})\{?(.+?)(?:\}|\\end\{keyword\})", src, re.S)
    if m:
        lines.append("Keywords: " + _strip_latex(m.group(1)).replace("\\sep", ","))
    body = src.split("\\begin{document}", 1)[-1]
    body = re.sub(r"\\begin\{abstract\}.+?\\end\{abstract\}", "", body, flags=re.S)
    for sec in re.split(r"\\(?:sub)*section\*?\{([^}]*)\}", body):
        lines.extend(_strip_latex(sec).splitlines())
    bib = re.findall(r"\\bibitem(?:\[[^\]]*\])?\{[^}]*\}(.+?)(?=\\bibitem|\\end\{thebibliography\})", src, re.S)
    if bib:
        lines.append("References")
        lines.extend(_strip_latex(b) for b in bib)
    lines.extend(f"Figure {i + 1}" for i in range(len(re.findall(r"\\begin\{figure\*?\}", src))))
    lines.extend(f"Table {i + 1}" for i in range(len(re.findall(r"\\begin\{table\*?\}", src))))
    return lines


def _strip_latex(s: str) -> str:
    s = re.sub(r"\\(?:emph|textbf|textit|mathrm|text)\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?(\{[^}]*\})?", " ", s)
    s = s.replace("{", "").replace("}", "").replace("~", " ")
    return re.sub(r"[ \t]+", " ", s).strip()


def parse(path: str | Path) -> Manuscript:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    lines = [l.strip() for l in read_text(path)]
    lines = [l for l in lines if l]
    ms = Manuscript(path=str(path))
    if not lines:
        return ms

    # Title: first reasonably long line that isn't a label like "Original article".
    title_at = -1
    for i, line in enumerate(lines[:15]):
        if 15 <= len(line) <= 300 and not HEADING_ABSTRACT_RE.match(line) \
                and not re.match(r"^(original|research|review)\s+(article|paper)$", line, re.I):
            ms.title = line.rstrip(".")
            title_at = i
            break

    # Abstract: paragraph(s) after an "Abstract" heading until the next section.
    ref_start = len(lines)
    for i, line in enumerate(lines):
        if HEADING_REF_RE.match(line):
            ref_start = i  # keep the last match: some papers mention "References" earlier
    for i, line in enumerate(lines):
        m = HEADING_ABSTRACT_RE.match(line)
        if m and i < ref_start:
            parts = [m.group(2)] if m.group(2) else []
            for nxt in lines[i + 1:i + 30]:
                if SECTION_AFTER_ABSTRACT_RE.match(nxt) or KEYWORDS_RE.match(nxt):
                    break
                parts.append(nxt)
                if sum(len(p) for p in parts) > 3000:
                    break
            ms.abstract = " ".join(p for p in parts if p).strip()
            if 0 <= title_at < i:
                ms.front_matter = "\n".join(lines[title_at + 1:i])[:4000]
            break

    for line in lines[:ref_start]:
        m = KEYWORDS_RE.match(line)
        if m:
            ms.keywords = [k.strip(" .") for k in re.split(r"[;,·•]|\s\|\s", m.group(2)) if k.strip(" .")]
            break

    body = lines[:ref_start]
    ms.declarations = extract_declarations(lines[:ref_start])
    ms.body_text = "\n".join(body)
    ms.word_count = sum(len(l.split()) for l in body)
    ms.references = _split_references(lines[ref_start + 1:]) if ref_start < len(lines) else []
    ms.reference_dois = sorted({_clean_doi(d) for d in DOI_RE.findall("\n".join(lines[ref_start:]))})
    ms.figure_count = len({int(n) for n in FIGURE_RE.findall(ms.body_text)})
    ms.table_count = len({int(n) for n in TABLE_RE.findall(ms.body_text)})
    return ms


def _clean_doi(doi: str) -> str:
    return doi.rstrip(".)]").lower()


def _split_references(lines: list[str]) -> list[str]:
    """Group reference lines into entries; PDF extraction wraps entries over lines."""
    refs: list[str] = []
    start_re = re.compile(r"^\s*(\[\d+\]|\d+\.\s|\d+\s|[A-Z][A-Za-z'’\-]+,\s*[A-Z])")
    for line in lines:
        if start_re.match(line) or not refs:
            refs.append(line)
        else:
            refs[-1] += " " + line
    return [r.strip() for r in refs if len(r) > 20]


DECLARATION_HEADINGS = {
    "competing_interests": r"(declaration\s+of\s+)?(competing|conflicts?\s+of)\s+interests?|conflicts?\s+of\s+interest|"
                           r"disclosure\s+statement|declaration\s+of\s+interests?",
    "funding": r"funding(\s+(information|statement|sources?|details))?|financial\s+(support|disclosure)|"
               r"sources?\s+of\s+funding|grant\s+information",
    "data_availability": r"data\s+(availability|access|sharing)(\s+statement)?|availability\s+of\s+(data|data\s+and\s+materials?)",
    "ethics": r"ethic(s|al)(\s+(approval|statement|declarations?|considerations?))?(\s+and\s+consent.*)?|"
              r"consent\s+to\s+participate|informed\s+consent",
    "author_contributions": r"authors?'?\s+contributions?|credit\s+(authorship\s+)?(contribution\s+)?statement|"
                            r"contributions?\s+of\s+the\s+authors",
    "preprint": r"preprint(\s+statement)?",
    "acknowledgements": r"acknowledge?ments?",
    "ai_disclosure": r"(declaration\s+of\s+)?(the\s+)?(use\s+of\s+)?(generative\s+)?(ai|artificial\s+intelligence)"
                     r"(\s+and\s+ai[\s-]+assisted\s+technologies)?(\s+(use|tools?|disclosure|statement|declaration))?(\s+in\s+"
                     r"(the\s+)?(writing|preparation)(\s+process)?)?|declaration\s+of\s+generative\s+ai.*",
}


def extract_declarations(lines: list[str]) -> dict:
    """Funding, competing interests, data availability, ethics, contributions and preprint
    statements written in the manuscript (usually under 'Declarations' before the references)."""
    found: dict[str, str] = {}
    current, buf = None, []

    def close():
        if current and buf:
            text = " ".join(buf).strip()
            if len(text) >= 2 and current not in found:
                found[current] = text[:800]
            elif len(text) >= 2 and text not in found[current]:  # e.g. "Informed consent" and "Ethics approval"
                found[current] = (found[current] + " " + text)[:1200]

    for line in lines:
        # "Funding: None.", "Funding. None." or a heading line on its own ("Conflict of interest").
        key, rest = None, ""
        splits = [line.split(":", 1)] if ":" in line[:80] else []
        splits += [[line[:m.start()], line[m.end():]] for m in re.finditer(r"\.\s+", line[:81])][:3]
        for head, tail in splits + [[line, ""]]:
            head = re.sub(r"^\s*\d+(\.\d+)*\.?\s*", "", head).strip(" .-–—*")
            if len(head) <= 80:
                key = next((k for k, rx in DECLARATION_HEADINGS.items() if re.fullmatch(rx, head, re.I)), None)
            if key:
                rest = tail
                break
        if key:
            close()
            current, buf = key, ([rest.strip()] if rest.strip() else [])
            continue
        if current:
            # A new short heading ends the statement.
            if len(line) < 60 and not line.endswith(".") and buf:
                close()
                current, buf = None, []
                continue
            buf.append(line.strip())
            if sum(len(b) for b in buf) > 800:
                close()
                current, buf = None, []
    close()
    return found
