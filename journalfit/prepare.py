"""Build a submission package for one journal: manuscript copy, cover letter,
guideline checklist and the metadata the submission forms ask for."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from .models import Candidate, Manuscript


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:60]


def build_package(ms: Manuscript, cand: Candidate, llm, author: dict, out_root: Path, http=None,
                  supplements: list[str] | None = None, declarations: dict | None = None,
                  authors: list[dict] | None = None, reviewers: list[dict] | None = None, log=None) -> Path:
    j = cand.journal
    out = out_root / slug(j.name)
    out.mkdir(parents=True, exist_ok=True)
    src = Path(ms.path)
    shutil.copy2(src, out / f"manuscript{src.suffix}")
    supp_dir = out / "supplementary"
    if supp_dir.exists():
        shutil.rmtree(supp_dir)  # the package always reflects the files given this time
    for f in supplements or []:
        supp_dir.mkdir(exist_ok=True)
        shutil.copy2(f, supp_dir / Path(f).name)

    guide: dict = {}
    if llm:
        try:
            page = ""
            if not llm.has_web:  # e.g. DeepSeek: find and read the guidelines page ourselves
                from .http import HttpClient
                from .metrics.publishers import PageFetcher, find_guidelines_text

                fetcher = PageFetcher(http or HttpClient(), remote=getattr(llm, "fetch_text", None))
                try:
                    url, page = find_guidelines_text(j, fetcher)
                finally:
                    fetcher.close()
            guide = llm.guidelines(ms, j, page)
            if page and not guide.get("guidelines_url"):
                guide["guidelines_url"] = url
        except Exception as e:
            guide = {"error": str(e)}
        letter = llm.cover_letter(ms, j, author, cand.scope_reason, guide.get("required_items") or [],
                                  declarations or {})
        _write_docx(out / "cover_letter.docx", letter)
        (out / "cover_letter.txt").write_text(letter, encoding="utf-8")

    say = log or (lambda m: None)
    declarations = dict(declarations or {})
    made = make_files(ms, out, guide, llm, author, declarations, authors, say)
    if reviewers is None:
        reviewers = suggest_reviewers(ms, cand, made["authors"], guide, http, say)
    from .submit import people

    people.save(out, "reviewers.json", reviewers)

    metadata = {
        "journal": j.name, "publisher": j.publisher, "issn": j.issns,
        "references": len(ms.references), "abstract_words": len(ms.abstract.split()),
        "funders": made["funders"], "peer_review": guide.get("peer_review") or "unknown",
        "generated_files": made["files"],
        "title": ms.title, "abstract": ms.abstract, "keywords": ms.keywords,
        "article_type": ms.article_type, "word_count": ms.word_count,
        "figures": ms.figure_count, "tables": ms.table_count,
        "corresponding_author": author,
        "submission_system": guide.get("submission_system"),
        "submission_url": guide.get("submission_url"),
        "guidelines_url": guide.get("guidelines_url"),
        "supplementary_files": sorted(p.name for p in supp_dir.glob("*")) if supp_dir.exists() else [],
        "declarations": declarations or {},
    }
    (out / "metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    (out / "checklist.md").write_text(_checklist(ms, cand, guide, metadata["supplementary_files"], declarations or {})
                                      + _ready_section(made, reviewers, guide, declarations), encoding="utf-8")
    return out


def make_files(ms: Manuscript, out: Path, guide: dict, llm, author: dict, declarations: dict,
               authors: list[dict] | None, say) -> dict:
    """Authors list, title page, anonymised copy, figures, highlights and funders for this journal."""
    from . import authors as authors_mod
    from . import config, files
    from .submit import people

    made: dict = {"files": [], "notes": [], "funders": []}
    src = Path(ms.path)
    found = authors if authors is not None else authors_mod.from_manuscript(str(src), ms.front_matter)
    if not found and authors is None and llm and ms.front_matter:
        try:
            found = llm.extract_authors(ms.front_matter)
        except Exception as e:
            made["notes"].append(f"Could not read the author list ({e}).")
    me = {"name": author.get("name", ""), "email": author.get("email", ""), "institution": author.get("affiliation", "")}
    book_path = config.DATA_DIR / "people.json"
    team = people.merge_authors(found, me, people.address_book(book_path))
    people.save(out, "authors.json", team)
    if found:
        people.remember(book_path, team)
    made["authors"] = team
    say(f"Authors: {', '.join(people.full_name(a) for a in team) or 'none found'}")

    funding = declarations.get("funding", "")
    if llm and funding:
        from .submit.assistant import statement_says_no

        if statement_says_no(funding) is False:
            try:
                made["funders"] = llm.funders(funding)
            except Exception:
                pass

    if team:
        # Statements written in the manuscript (e.g. acknowledgements) fill what the author didn't give here.
        files.title_page(out / "title_page.docx", ms.title, team, {**ms.declarations, **declarations}, ms.word_count)
        made["files"].append("title_page.docx")
    blind = (guide.get("peer_review") or "").lower().startswith("double")
    for old in out.glob("anonymised_manuscript.*"):
        old.unlink()
    if blind:
        if src.suffix.lower() == ".docx":
            removed = files.anonymise(src, out / "anonymised_manuscript.docx", team)
            made["files"].append("anonymised_manuscript.docx")
            made["notes"].append("Double-blind review: the anonymised copy is uploaded as the manuscript. Removed: "
                                 + "; ".join(removed[:8]))
            say("Made an anonymised manuscript for double-blind review.")
        else:
            made["notes"].append("Double-blind review: remove author names and affiliations from the manuscript "
                                 "yourself (automatic anonymising works on Word files).")
    fig_dir = out / "figures"
    if fig_dir.exists():
        shutil.rmtree(fig_dir)
    if guide.get("figures_separate") and src.suffix.lower() == ".docx":
        figs = files.extract_figures(src, fig_dir)
        if figs:
            made["files"] += [f"figures/{f.name}" for f in figs]
            made["notes"].append("Separate figure files: " + "; ".join(files.figure_sizes(figs))
                                 + ". Print quality is usually at least 300 dpi at the printed size.")
    for old in ("highlights.txt", "highlights.docx"):
        (out / old).unlink(missing_ok=True)
    if guide.get("highlights_required") and llm and ms.abstract:
        try:
            bullets = llm.highlights(ms, int(guide.get("highlights_max_chars") or 85))
        except Exception:
            bullets = []
        if bullets:
            files.highlights_files(out, bullets)
            made["files"].append("highlights.docx")
            made["notes"].append("Highlights were written by the AI from your abstract: read them before submitting.")
    return made


def suggest_reviewers(ms: Manuscript, cand: Candidate, team: list[dict], guide: dict, http, say) -> list[dict]:
    from . import reviewers as reviewers_mod
    from .http import HttpClient

    queries = (ms.search_queries or [])[:3] or [ms.title]
    want = max(int(guide.get("suggested_reviewers") or 0), 5)
    try:
        found = reviewers_mod.suggest(http or HttpClient(), queries, team, n=want, journal_name=cand.journal.name)
    except Exception as e:
        say(f"Could not suggest reviewers ({e}).")
        return []
    say(f"Suggested reviewers: {len(found)} (add their emails before submitting).")
    return found


def _ready_section(made: dict, reviewers: list[dict], guide: dict, declarations: dict) -> str:
    lines = ["", "## Prepared for you"]
    lines += [f"- {f}" for f in made["files"]] or ["- (no extra files needed)"]
    lines += [f"- {n}" for n in made["notes"]]
    todo = []
    team = made["authors"]
    if not team:
        todo.append("Add the authors (none could be read from the manuscript).")
    for a in team:
        gaps = [k for k in ("email", "institution", "country") if not a.get(k)]
        if gaps:
            todo.append(f"{' '.join(x for x in (a.get('given'), a.get('family')) if x)}: add {', '.join(gaps)}")
    no_email = [r for r in reviewers if not r.get("email")]
    if no_email:
        todo.append(f"Suggested reviewers: add emails for {len(no_email)} (they are typed into the form only "
                    "with an email). Check each has no conflict of interest with you.")
    if not declarations.get("ai_disclosure"):
        todo.append("Use of generative AI: say whether you used AI tools (many journals now ask).")
    if todo:
        lines += ["", "## Still needed from you"] + [f"- [ ] {t}" for t in todo]
    return "\n".join(lines) + "\n"


def _write_docx(path: Path, text: str) -> None:
    import docx

    d = docx.Document()
    for para in text.split("\n\n"):
        d.add_paragraph(para.strip())
    d.save(str(path))


def _checklist(ms: Manuscript, cand: Candidate, guide: dict, supplements: list[str] = (),
               declarations: dict | None = None) -> str:
    j = cand.journal
    lines = [f"# Submission checklist — {j.name}", ""]
    if guide.get("error"):
        lines += [f"> Could not read the author guidelines automatically: {guide['error']}",
                  "> Check them on the journal website before submitting.", ""]
    if guide.get("guidelines_url"):
        lines.append(f"Author guidelines: {guide['guidelines_url']}")
    if guide.get("submission_url"):
        lines.append(f"Submission site ({guide.get('submission_system') or 'unknown system'}): {guide['submission_url']}")
    if guide.get("format_free_first_submission"):
        lines += ["", "Good news: this journal accepts any reasonable format and reference style for the "
                      "first submission. Formatting can wait until revision."]
    lines.append("")
    if guide.get("article_type_accepted") is False:
        lines += [f"**Warning:** the guidelines suggest this journal does not accept '{ms.article_type}'.", ""]

    issues = guide.get("issues") or []
    items = guide.get("required_items") or ["Competing interests statement", "Funding statement",
                                            "Data availability statement", "Author contributions"]
    lines.append("## Must do before submitting")
    lines += [f"- [ ] **{i.get('item')}**: yours = {i.get('manuscript')}; required = {i.get('requirement')}"
              for i in issues]
    lines += [f"- [ ] {item}" for item in items]
    lines += ["- [ ] All co-authors have approved this version and the journal choice",
              "- [ ] The manuscript is not under consideration at any other journal",
              "- [ ] Cover letter read, and [placeholders] filled in"]
    if declarations:
        from .settings import DECLARATION_FIELDS

        lines += ["", "## Your declarations (used in the cover letter and the submission forms)"]
        lines += [f"- {DECLARATION_FIELDS.get(k, k)}: {v}" for k, v in declarations.items() if k in DECLARATION_FIELDS]
        if declarations.get("confirm_standard"):
            lines.append("- Standard author confirmations: you asked the tool to tick these on the journal's site")
    else:
        lines += ["", "Tip: add your declarations (competing interests, funding, data availability...) "
                      "in Settings and they'll be filled in for you."]
    if supplements:
        lines += ["", "## Supplementary files in this package"] + [f"- {n}" for n in supplements]

    lines += ["", "## Facts"]
    for k, label in (("word_limit", "Word limit"), ("abstract_word_limit", "Abstract limit"),
                     ("keyword_limit", "Keywords"), ("reference_style", "Reference style")):
        if guide.get(k):
            lines.append(f"- {label}: {guide[k]}")
    lines.append(f"- Your manuscript: {ms.word_count} words, abstract {len(ms.abstract.split())} words, "
                 f"{len(ms.keywords)} keywords, {ms.figure_count} figures, {ms.table_count} tables")

    minor = guide.get("minor_issues") or []
    later = guide.get("can_wait") or []
    if minor or later:
        lines += ["", "## Can wait (formatting; usually only needed at revision or acceptance)"]
        lines += [f"- {i.get('item')}: yours = {i.get('manuscript')}; journal style = {i.get('requirement')}" for i in minor]
        lines += [f"- {x}" for x in later]
    return "\n".join(lines) + "\n"
