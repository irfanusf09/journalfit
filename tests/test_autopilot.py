"""The autopilot on a multi-step practice submission site (tests/portal): it fills every
step, adds co-authors, answers the Yes/No questions from the author's statements, uploads the
files with their item types, and stops at the review page without pressing Submit."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from journalfit.submit import assistant, navigator, people

playwright = pytest.importorskip("playwright.sync_api")

PORTAL = Path(__file__).parent / "portal"
AUTHORS = [
    {"given": "Irfan", "family": "Yousuf", "email": "irfan@uni.example", "institution": "University of Kashmir",
     "country": "India", "corresponding": True},
    {"given": "Asha", "family": "Rao", "email": "asha.rao@uni.example", "institution": "IIT Delhi",
     "country": "India", "orcid": "0000-0002-1825-0097"},
    {"given": "John", "family": "Smith", "email": "j.smith@lse.example", "institution": "LSE", "country": "UK"},
]
REVIEWERS = [{"given": "Maria", "family": "Garcia", "email": "m.garcia@uv.example", "institution": "Univ. of Valencia",
              "reason": "Works on household saving behaviour."}]


def make_package(tmp_path: Path) -> Path:
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    (pkg / "manuscript.docx").write_bytes(b"x")
    (pkg / "cover_letter.docx").write_bytes(b"x")
    (pkg / "cover_letter.txt").write_text("Dear Editor, please consider our paper.")
    (pkg / "title_page.docx").write_bytes(b"x")
    (pkg / "figures").mkdir()
    for n in (2, 1, 10):
        (pkg / "figures" / f"Figure_{n}.png").write_bytes(b"x")
    (pkg / "metadata.json").write_text(json.dumps({
        "journal": "J", "title": "Five Dimensions of Everyday Financial Behaviour in Kashmir",
        "abstract": "We study saving.", "keywords": ["saving", "credit"], "article_type": "original research",
        "word_count": 7200, "figures": 3, "tables": 2,
        "declarations": {
            "competing_interests": "The authors declare no competing interests.",
            "funding": "This research did not receive any specific grant.",
            "data_availability": "Data are available on request.",
            "previous_submission": "This manuscript has not been submitted to this journal before.",
        }}))
    people.save(pkg, "authors.json", AUTHORS)
    people.save(pkg, "reviewers.json", REVIEWERS)
    return pkg


def run_portal(pkg: Path, monkeypatch, max_ticks: int = 120):
    from playwright.sync_api import sync_playwright

    monkeypatch.setattr(assistant, "UPLOAD_GAP", 0)
    logs: list[str] = []
    plan, _ = assistant.load_plan(pkg)
    driver = assistant.Driver(pkg, plan, log=logs.append)
    driver.pilot.AFTER_CLICK_WAIT = 1.5
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        page = b.new_page()
        page.goto((PORTAL / "step1.html").as_uri())
        page.evaluate("localStorage.clear()")
        page.reload()
        for _ in range(max_ticks):
            status = driver.tick(page)
            if status == "paused" and "review" in page.url:
                break
            page.wait_for_timeout(250)
        store = page.evaluate("Object.fromEntries(Object.entries(localStorage))")
        final_url = page.url
        b.close()
    return store, final_url, logs


def test_autopilot_reaches_review_page_and_never_submits(tmp_path, monkeypatch):
    pkg = make_package(tmp_path)
    store, url, logs = run_portal(pkg, monkeypatch)
    text = "\n".join(logs)
    assert url.endswith("review.html"), text
    # Never pressed: Submit, Approve, Build PDF, Cancel, Delete, Browse.
    for flag in ("submitted", "pdf", "cancelled", "deleted", "browse"):
        assert flag not in store, (flag, text)
    assert "AUTOPILOT PAUSED: this is the review / approval page" in text
    # Step 1
    assert store["v:type"] == "Original Article"
    assert store["v:title"].startswith("Five Dimensions")
    assert store["v:keywords"] == "saving; credit"
    # Step 2: the corresponding author was already listed; the two co-authors were added.
    assert (store["v:given1"], store["v:family1"], store["v:email1"]) == ("Asha", "Rao", "asha.rao@uni.example")
    assert store["v:orcid1"] == "0000-0002-1825-0097" and store["v:country1"] == "India"
    assert (store["v:given2"], store["v:family2"], store["v:country2"]) == ("John", "Smith", "United Kingdom")
    # Step 3: suggested reviewer filled, opposed reviewer boxes left alone.
    assert (store["v:r1given"], store["v:r1email"]) == ("Maria", "m.garcia@uv.example")
    assert store["v:r1reason"].startswith("Works on")
    assert store["v:opgiven"] == "" and store["v:opemail"] == ""
    # Step 4: numbers, Yes/No questions from the author's statements, statements, no newsletter.
    assert (store["v:words"], store["v:figs"], store["v:tabs"]) == ("7200", "3", "2")
    assert store["v:prev"] == "n" and store["v:fund"] == "n"
    assert store["v:coi"].startswith("The authors declare no")
    assert store["v:newsletter"] == ""
    # Step 5: every file with its own item type, figures in number order.
    assert store["files"].split("|") == [
        "Manuscript:manuscript.docx", "Title Page:title_page.docx", "Cover Letter:cover_letter.docx",
        "Figure:Figure_1.png", "Figure:Figure_2.png", "Figure:Figure_10.png"]
    score = assistant.automation_score(pkg)
    assert score["percent"] is not None and score["percent"] >= 90, score


def test_autopilot_pauses_when_a_required_box_cannot_be_filled(tmp_path, monkeypatch):
    pkg = make_package(tmp_path)
    meta = json.loads((pkg / "metadata.json").read_text())
    meta["declarations"].pop("previous_submission")  # the author hasn't said
    (pkg / "metadata.json").write_text(json.dumps(meta))
    store, url, logs = run_portal(pkg, monkeypatch, max_ticks=80)
    assert url.endswith("step4.html")
    assert any("AUTOPILOT PAUSED: fill these required boxes" in line and "submitted previously" in line
               for line in logs), logs
    assert "submitted" not in store


def test_forbidden_buttons_are_never_next():
    def b(text, **kw):
        return dict({"tag": "button", "type": "", "text": text, "words": [text], "href": "", "target": ""}, **kw)

    for t in ("Submit", "Build PDF for Approval", "Approve Submission", "I agree and continue", "Continue to payment",
              "Submit and Continue", "Finish", "Delete", "Cancel", "Back", "Save and Close", "Upload", ">",
              "Confirm", "Withdraw", "Transfer to another journal", "Next page of results"):
        assert navigator.classify(b(t)) is None, t
    for t in ("Next", "Save & Continue", "Continue", "Proceed", "Next Step", "Save and Next >"):
        assert navigator.classify(b(t)) == "next", t
    assert navigator.classify(b("Next", words=["Next", "Submit manuscript"])) is None  # hidden wording counts
    assert navigator.classify(b("Next", tag="a", href="https://elsewhere.example/")) is None
    assert navigator.classify(b("Add Co-Author")) == "add_author"
    assert navigator.classify(b("Add to My Authors")) == "add_author"
    assert navigator.classify(b("Add Reviewer")) == "add_reviewer"


def test_people_records():
    assert people.orcid_ok("0000-0002-1825-0097") and not people.orcid_ok("0000-0002-1825-0098")
    p = people.normalise({"name": "Dr. Irfan Yousuf", "orcid": "https://orcid.org/0000-0002-1825-0097",
                          "email": "not an email"})
    assert (p["given"], p["family"], p["orcid"], p["email"]) == ("Irfan", "Yousuf", "0000-0002-1825-0097", "")
    assert people.split_name("Yousuf, Irfan") == ("Irfan", "Yousuf")
    assert people.pick_option(["Select", "India", "United States"], "country", "USA") == "United States"
    assert people.pick_option(["Select", "Dr.", "Prof."], "salutation", "Professor") == "Prof."
    assert people.attribute(["Funding institution"]) == "institution"
    assert not people.is_person_field(["Funding institution"])  # funders are not people
    assert people.role(["First Name", "Opposed Reviewers"]) is None


def test_merge_authors_fills_gaps_from_settings_and_address_book():
    found = [{"name": "Irfan Yousuf"}, {"name": "Asha Rao", "email": "asha@x.example"}]
    me = {"name": "Irfan Yousuf", "email": "irfan@x.example", "institution": "Univ"}
    book = {"asha@x.example": people.normalise({"name": "Asha Rao", "orcid": "0000-0002-1825-0097"})}
    out = people.merge_authors(found, me, book)
    assert [people.full_name(p) for p in out] == ["Irfan Yousuf", "Asha Rao"]
    assert out[0]["corresponding"] and out[0]["email"] == "irfan@x.example"
    assert out[1]["orcid"] == "0000-0002-1825-0097"


def test_reviewer_suggestions_skip_authors_their_coauthors_and_weak_records(monkeypatch):
    from journalfit import reviewers

    def au(aid, name, inst, country="DE", pos="middle"):
        return {"author": {"id": f"https://openalex.org/{aid}", "display_name": name, "orcid": None},
                "institutions": [{"display_name": inst}], "countries": [country], "author_position": pos}

    works = [
        {"id": "W1", "title": "Saving nudges", "publication_year": 2025, "doi": "",
         "primary_location": {"source": {"display_name": "J Econ", "issn_l": "1234-5678"}},
         "authorships": [au("A1", "Anna Becker", "Univ Bonn", pos="last"), au("A2", "Carl Diaz", "Univ Bonn")]},
        {"id": "W2", "title": "Impulse buying", "publication_year": 2024, "doi": "",
         "primary_location": {"source": {"display_name": "J Psych", "issn_l": "1234-5678"}},
         "authorships": [au("A3", "Irfan Yousf", "LPU", "IN"), au("A4", "Eve Fox", "Oxford", "GB")]},
        {"id": "W3", "title": "Peer saving", "publication_year": 2024, "doi": "",
         "primary_location": {"source": {"display_name": "J Dev", "issn_l": "1234-5678"}},
         "authorships": [au("A5", "Gita Rao", "LPU", "IN"), au("A6", "Hugo Lind", "Lund", "SE"),
                         au("A7", "Ian Moe", "Oslo", "NO")]},
    ]

    class Http:
        def get_json(self, url, params):
            if url.endswith("/works"):
                return {"results": works}
            return {"results": [{"id": f"https://openalex.org/A{i}", "works_count": 100 if i != 7 else 3,
                                 "summary_stats": {"h_index": 20}} for i in range(1, 8)]}

    monkeypatch.setattr(reviewers, "_indexed_issns", lambda: {"12345678"})
    team = [{"given": "Irfan", "family": "Yousf", "institution": "LPU"}]
    got = [f"{r['given']} {r['family']}" for r in reviewers.suggest(Http(), ["q"], team, n=5)]
    # Not: the author; Eve Fox (co-author of the author); Gita Rao (author's institution); Ian Moe (3 papers);
    # Carl Diaz (same institution and paper as Anna Becker, who is picked first).
    assert got == ["Anna Becker", "Hugo Lind"], got


def test_package_has_authors_title_page_blinded_copy_highlights_and_reviewers(tmp_path, monkeypatch):
    import docx

    from journalfit import manuscript, prepare
    from journalfit.models import Candidate, Journal

    src = tmp_path / "paper.docx"
    d = docx.Document()
    d.add_paragraph("Saving Mechanisms and Peers in Kashmir Households")
    p = d.add_paragraph("Irfan Yousf")
    p.add_run("1*").font.superscript = True
    p.add_run(", Asha Rao")
    p.add_run("2").font.superscript = True
    q = d.add_paragraph()
    q.add_run("1").font.superscript = True
    q.add_run(" Lovely Professional University, Phagwara, Punjab, India")
    q = d.add_paragraph()
    q.add_run("2").font.superscript = True
    q.add_run(" Department of Economics, University of Delhi, Delhi, India")
    d.add_paragraph("*Corresponding author: Irfan Yousf, irfan@lpu.example")
    d.add_paragraph("Abstract")
    d.add_paragraph("We survey 600 households and find that peers shape saving. " * 3)
    d.add_paragraph("Keywords: saving; peers")
    d.add_paragraph("1. Introduction")
    d.add_paragraph("Our team at Lovely Professional University ran the survey.")
    d.add_paragraph("Acknowledgements")
    d.add_paragraph("We thank Prof. X of the University of Delhi.")
    d.add_paragraph("References")
    d.add_paragraph("Yousf, I. (2024). Earlier work. Journal of Saving, 1, 1-2.")
    d.core_properties.author = "Irfan Yousf"
    d.save(str(src))
    ms = manuscript.parse(src)

    class LLM:
        has_web = True

        def guidelines(self, ms, j, page=""):
            return {"peer_review": "double-blind", "highlights_required": True, "figures_separate": False,
                    "submission_url": "https://x.example/submit"}

        def cover_letter(self, *a, **k):
            return "Dear Editor"

        def highlights(self, ms, max_chars=85):
            return ["Peers shape household saving in Kashmir", "A survey of 600 households"]

        def funders(self, statement):
            return [{"name": "Indian Council of Social Science Research", "award": "ICSSR-123"}]

    monkeypatch.setattr(prepare, "suggest_reviewers", lambda *a: [
        {"given": "Anna", "family": "Becker", "email": "", "institution": "Bonn", "reason": "Saving nudges"}])
    cand = Candidate(journal=Journal(openalex_id="S1", name="Journal of Saving"))
    out = prepare.build_package(ms, cand, LLM(), {"name": "Irfan Yousf", "email": "irfan@lpu.example"}, tmp_path / "out",
                                declarations={"funding": "Funded by the Indian Council of Social Science Research "
                                                         "(ICSSR-123).", "competing_interests": "None."})
    team = json.loads((out / "authors.json").read_text())
    assert [(a["given"], a["family"]) for a in team] == [("Irfan", "Yousf"), ("Asha", "Rao")]
    assert team[0]["corresponding"] and team[0]["email"] == "irfan@lpu.example"
    assert team[1]["department"] == "Department of Economics" and team[1]["institution"] == "University of Delhi"
    blinded = docx.Document(str(out / "anonymised_manuscript.docx"))
    text = "\n".join(p.text for p in blinded.paragraphs)
    body = text.split("References")[0]
    assert "Yousf" not in body and "Lovely Professional" not in body and "Asha" not in body
    assert "Acknowledgements" not in body and "Prof. X" not in body
    assert "Yousf, I. (2024)" in text  # references are left as they are
    assert blinded.core_properties.author == ""
    title_page = "\n".join(p.text for p in docx.Document(str(out / "title_page.docx")).paragraphs)
    assert "Irfan Yousf" in title_page and "University of Delhi" in title_page and "We thank" in title_page
    assert (out / "highlights.txt").read_text().startswith("Peers shape")
    plan, meta = assistant.load_plan(out)
    assert plan.files["manuscript_file"].endswith("anonymised_manuscript.docx")
    assert plan.files["title_page_file"].endswith("title_page.docx")
    assert plan.values["funder_name"] == "Indian Council of Social Science Research"
    assert plan.values["grant_number"] == "ICSSR-123"
    assert plan.people["reviewer"] == []  # no email yet, so never typed into a form
    checklist = (out / "checklist.md").read_text()
    assert "add emails for 1" in checklist and "Use of generative AI" in checklist
    assert json.loads((tmp_path / "out" / "journal-of-saving" / "authors.json").read_text())
