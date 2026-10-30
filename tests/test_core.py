"""Offline tests: parsing, metric extraction, indexing data, filters and the full
finder pipeline against a stubbed OpenAlex/publisher web."""

from __future__ import annotations

import json

import docx
import pytest

from journalfit import manuscript
from journalfit.config import Criteria
from journalfit.finder import Finder, final_score
from journalfit.http import FetchError
from journalfit.metrics.extract import extract_metrics
from journalfit.metrics.publishers import candidate_urls
from journalfit.models import Candidate, Journal, Metrics
from journalfit.sources.indexing import IndexingData, parse_scimago

ABSTRACT = ("Lithium-sulfur batteries suffer from polysulfide shuttling. We design a graphene oxide "
            "interlayer doped with nitrogen that traps polysulfides and improves cycling stability "
            "of lithium sulfur cells over 500 cycles at high sulfur loading.")


def make_docx(path):
    d = docx.Document()
    d.add_paragraph("Original article")
    d.add_paragraph("Nitrogen-doped graphene oxide interlayers suppress polysulfide shuttling in lithium-sulfur batteries")
    d.add_paragraph("A. Author, B. Author")
    d.add_paragraph("Abstract")
    d.add_paragraph(ABSTRACT)
    d.add_paragraph("Keywords: lithium-sulfur battery; graphene oxide; interlayer; polysulfide")
    d.add_paragraph("1. Introduction")
    d.add_paragraph("Batteries matter. As shown in Figure 1 and Figure 2 and Table 1, results improve.")
    d.add_paragraph("References")
    d.add_paragraph("[1] Smith J, et al. A paper. Carbon 2020. https://doi.org/10.1016/j.carbon.2020.01.001")
    d.add_paragraph("[2] Doe K. Another paper on sulfur cathodes. J Mater Sci 2021. doi:10.1007/s10853-021-0001-2.")
    d.save(path)


def test_parse_docx(tmp_path):
    p = tmp_path / "paper.docx"
    make_docx(p)
    ms = manuscript.parse(p)
    assert ms.title.startswith("Nitrogen-doped graphene oxide")
    assert ms.abstract == ABSTRACT
    assert ms.keywords == ["lithium-sulfur battery", "graphene oxide", "interlayer", "polysulfide"]
    assert ms.reference_dois == ["10.1007/s10853-021-0001-2", "10.1016/j.carbon.2020.01.001"]
    assert len(ms.references) == 2
    assert (ms.figure_count, ms.table_count) == (2, 1)


def test_parse_latex(tmp_path):
    p = tmp_path / "paper.tex"
    p.write_text(r"""
\documentclass{article}
\title{Deep learning for crop yield prediction from satellite images}
\begin{document}
\maketitle
\begin{abstract}
We predict \emph{crop yield} using convolutional networks on Sentinel-2 imagery across 300 districts.
\end{abstract}
\keywords{remote sensing, crop yield, deep learning}
\section{Introduction}
Text here.
\begin{thebibliography}{9}
\bibitem{a} Author A. Paper about yield. Remote Sens. 2020. doi:10.3390/rs12010001
\end{thebibliography}
\end{document}
""")
    ms = manuscript.parse(p)
    assert ms.title == "Deep learning for crop yield prediction from satellite images"
    assert "crop yield using convolutional networks" in ms.abstract
    assert ms.keywords == ["remote sensing", "crop yield", "deep learning"]
    assert ms.reference_dois == ["10.3390/rs12010001"]


@pytest.mark.parametrize("text,field,expected", [
    ("Submission to first decision (median) 21 days Downloads 2023", "days_to_first_decision", 21),
    ("Review time 6.4 weeks Submission to acceptance 14.2 weeks", "days_to_first_decision", 44.8),
    ("Review time 6.4 weeks Submission to acceptance 14.2 weeks", "days_to_acceptance", 99.4),
    ("Acceptance rate 18% Impact factor 9.4", "acceptance_rate", 18),
    ("Acceptance rate 18% Impact factor 9.4", "impact_factor", 9.4),
    ("5-year impact factor 7.0 CiteScore 15.2", "impact_factor", None),
    ("5-year impact factor 7.0 CiteScore 15.2", "citescore", 15.2),
    ("First decision: 17.3 days (median values for papers published in this journal)", "days_to_first_decision", 17.3),
    ("Time to first decision 2 months", "days_to_first_decision", 60.8),
    ("Our acceptance rate is low", "acceptance_rate", None),
    ("CiteScore 2025: 9.8 Downloads 120,000", "citescore", 9.8),
    ("Journal Impact Factor (2024) 5.2 Submission to first decision (median) 9 days", "impact_factor", 5.2),
    ("Impact Factor 2024 6.3", "impact_factor", 6.3),
    ("Impact factor 4.9 (2024)", "impact_factor", 4.9),
])
def test_extract_metrics(text, field, expected):
    assert getattr(extract_metrics(text, "u"), field) == expected


SCIMAGO = """Rank;Sourceid;Title;Type;Issn;SJR;SJR Best Quartile;H index;Publisher;Categories
1;1;Carbon;journal;"00086223, 18733891";2,171;Q1;300;Elsevier Ltd;Materials Science (miscellaneous) (Q1); Chemistry (Q1)
2;2;Journal of Materials Science;journal;"00222461, 15734803";0,8;Q2;200;Springer;Materials Science (Q2)
3;3;Some Proceedings;conference and proceedings;12345678;0,2;Q4;10;X;Y
"""


def test_scimago_parse_and_annotate():
    data = parse_scimago(SCIMAGO)
    assert data["00086223"].quartile == "Q1" and data["18733891"].sjr == 2.171
    assert "12345678" not in data
    idx = IndexingData()
    idx.scimago = data
    j = Journal("S1", "Carbon", issns=["0008-6223"])
    idx.annotate(j)
    assert j.scopus_indexed is True and j.metrics.quartile == "Q1" and j.metrics.h_index == 300
    j2 = Journal("S9", "Unknown", issns=["1111-2222"])
    idx.annotate(j2)
    assert j2.scopus_indexed is False


def test_metric_evidence_is_recorded():
    ev = {}
    extract_metrics("Journal metrics. Submission to first decision (median) 21 days. More text.", "u", ev)
    assert ev["days_to_first_decision"] == "Submission to first decision (median) 21 days"


def test_publisher_urls():
    j = Journal("S1", "Carbon", publisher="Elsevier BV", homepage="https://www.sciencedirect.com/journal/carbon")
    assert candidate_urls(j)[0] == "https://www.sciencedirect.com/journal/carbon/about/insights"
    j = Journal("S2", "J Mater Sci", publisher="Springer Science+Business Media",
                homepage="http://link.springer.com/journal/10853")
    assert "https://link.springer.com/journal/10853" in candidate_urls(j)
    j = Journal("S3", "Mat Res Lett", publisher="Taylor & Francis", homepage="https://www.tandfonline.com/toc/tmrl20/current")
    assert "journalCode=tmrl20" in candidate_urls(j)[0]
    j = Journal("S4", "Nature Human Behaviour", publisher="Nature Portfolio", homepage="https://www.nature.com/nathumbehav/")
    assert candidate_urls(j)[0] == "https://www.nature.com/nathumbehav/journal-impact"


def _cand(**m):
    c = Candidate(journal=Journal("S1", "J", metrics=Metrics(**m)), text_fit=0.8)
    return c


def test_metric_filters_and_missing_policy():
    f = Finder.__new__(Finder)
    crit = Criteria(max_first_decision_days=90, min_acceptance_rate=20)
    assert f.apply_metric_filters(_cand(days_to_first_decision=120, acceptance_rate=30), crit).startswith("first decision")
    assert f.apply_metric_filters(_cand(days_to_first_decision=40, acceptance_rate=10), crit).startswith("acceptance rate")
    keep = _cand(days_to_first_decision=40)
    assert f.apply_metric_filters(keep, crit) == "" and keep.warnings == ["acceptance rate not published"]
    crit.if_missing = "skip"
    assert f.apply_metric_filters(_cand(days_to_first_decision=40), crit) == "acceptance rate not published"


def test_faster_acceptance_scores_higher():
    fast = Candidate(journal=Journal("a", "A", metrics=Metrics(quartile="Q1", days_to_acceptance=90)), text_fit=0.8)
    slow = Candidate(journal=Journal("b", "B", metrics=Metrics(quartile="Q1", days_to_acceptance=330)), text_fit=0.8)
    assert final_score(fast, 1) > final_score(slow, 1)


def test_poor_scope_fit_is_excluded(tmp_path, monkeypatch):
    monkeypatch.setattr("journalfit.config.OVERRIDES_PATH", tmp_path / "none.csv")
    monkeypatch.setattr("journalfit.config.DATA_DIR", tmp_path)

    class ScopeLLM:
        def review_scope(self, ms, cands):
            for c in cands:
                c.scope_fit = 0.2 if c.journal.name == "Journal of Materials Science" else 0.9

        def lookup_metrics(self, j, pages=None):
            return Metrics()

    p = tmp_path / "paper.docx"
    make_docx(p)
    idx = IndexingData()
    idx.scimago = parse_scimago(SCIMAGO)
    f = Finder(FakeHttp(), llm=ScopeLLM(), indexing=idx, use_browser=False, log=lambda *_: None)
    res = f.find(manuscript.parse(p), Criteria(scopus=True), top=5)
    assert "Journal of Materials Science" not in [c.journal.name for c in res.ranked]
    reasons = {c.journal.name: c.excluded_reason for c in res.excluded}
    assert reasons["Journal of Materials Science"] == "poor scope fit (2/10): desk rejection likely"
    assert any("Only" in n for n in res.notes)


# --------------------------------------------------------------------------- end to end

def _src(sid, name, publisher, issn, homepage, is_oa=False):
    return {"id": f"https://openalex.org/{sid}", "type": "journal", "display_name": name,
            "host_organization_name": publisher, "issn_l": issn, "issn": [issn],
            "homepage_url": homepage, "is_oa": is_oa, "is_in_doaj": is_oa, "apc_usd": 2500 if is_oa else None,
            "summary_stats": {"h_index": 100}}


SOURCES = {
    "S1": _src("S1", "Carbon", "Elsevier BV", "0008-6223", "https://www.sciencedirect.com/journal/carbon"),
    "S2": _src("S2", "Journal of Materials Science", "Springer", "0022-2461", "http://link.springer.com/journal/10853"),
    "S3": _src("S3", "Batteries", "MDPI", "2313-0105", "https://www.mdpi.com/journal/batteries", is_oa=True),
    "S4": _src("S4", "Slow Journal", "Springer", "1234-5679", "http://link.springer.com/journal/99999"),
}
FILLER = "<p>" + "About the journal, aims and scope, editorial board and publishing information. " * 25 + "</p>"


def _page(name, body):
    return f"<html><body><h1>{name}</h1>{FILLER}{body}</body></html>"


PAGES = {
    "https://www.sciencedirect.com/journal/carbon/about/insights":
        _page("Carbon", "<p>Review time 4 weeks</p><p>Acceptance rate 25%</p>"),
    "https://link.springer.com/journal/10853":
        _page("Journal of Materials Science", "<div>Submission to first decision (median) 30 days</div>"),
    "https://link.springer.com/journal/99999":
        _page("Slow Journal", "<div>Submission to first decision (median) 200 days</div>"),
}


def _work(i, title, abstract, sid):
    return {"id": f"https://openalex.org/W{i}", "title": title, "publication_year": 2024,
            "abstract_inverted_index": {w: [k] for k, w in enumerate(abstract.split())},
            "primary_location": {"source": {"id": f"https://openalex.org/{sid}", "display_name": SOURCES[sid]["display_name"]}}}


WORKS = [
    _work(1, "Graphene interlayer for lithium sulfur batteries", "polysulfide shuttling graphene oxide interlayer lithium sulfur", "S1"),
    _work(2, "Nitrogen doped carbon hosts for sulfur cathodes", "nitrogen doped graphene polysulfide lithium sulfur cycling", "S1"),
    _work(3, "Sulfur cathode interlayers review", "lithium sulfur batteries interlayer polysulfide", "S2"),
    _work(4, "Polysulfide trapping in Li-S cells", "graphene oxide lithium sulfur polysulfide trapping", "S3"),
    _work(5, "Graphene oxide lithium sulfur interlayer cycling", "graphene oxide interlayer lithium sulfur", "S4"),
    _work(6, "Steel fatigue", "steel fatigue crack growth", "S2"),
]


CROSSREF_ITEMS = [
    {"DOI": f"10.1/x{i}", "assertion": [
        {"name": "received", "value": f"{1 + i % 20} January 2024"},
        {"name": "accepted", "value": f"{1 + i % 20} April 2024"},  # ~91 days later
        {"name": "first_online", "value": "1 May 2024"}]}
    for i in range(20)
] + [{"DOI": "10.1/no-dates", "assertion": [{"name": "copyright", "value": "Springer"}]}]


class FakeHttp:
    def get_json(self, url, params=None, ttl=None):
        if url.startswith("https://api.crossref.org/"):
            items = CROSSREF_ITEMS if "/journals/0022-2461/" in url else []
            return {"message": {"items": items}}
        if url.endswith("/works"):
            if params["filter"].startswith("doi:"):
                return {"results": [WORKS[2]]}
            return {"results": WORKS}
        if "/sources/" in url:
            return SOURCES[url.rsplit("/", 1)[-1]]
        if url.endswith("/publishers"):
            return {"results": [{"id": "https://openalex.org/P1", "display_name": "Springer", "works_count": 10}]}
        raise AssertionError(url)

    def get_text(self, url, params=None, ttl=None, headers=None, timeout=30):
        if url in PAGES:
            return PAGES[url]
        raise FetchError(url)


def test_finder_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr("journalfit.config.OVERRIDES_PATH", tmp_path / "none.csv")
    p = tmp_path / "paper.docx"
    make_docx(p)
    ms = manuscript.parse(p)
    idx = IndexingData()
    idx.scimago = parse_scimago(SCIMAGO + '4;4;Slow Journal;journal;"12345679";1;Q1;50;Springer;X\n')
    finder = Finder(FakeHttp(), llm=None, indexing=idx, use_browser=False, log=lambda *_: None)
    crit = Criteria(no_oa=True, scopus=True, max_first_decision_days=90, min_quartile="Q2")
    res = finder.find(ms, crit, top=5)

    names = [c.journal.name for c in res.ranked]
    assert names[0] == "Carbon"  # best topical match, fast, Q1, cited in references
    assert "Journal of Materials Science" in names
    reasons = {c.journal.name: c.excluded_reason for c in res.excluded}
    assert reasons["Batteries"] == "fully open access"
    assert reasons["Slow Journal"].startswith("first decision 200")
    jms = next(c for c in res.ranked if c.journal.name == "Journal of Materials Science").journal.metrics
    assert jms.days_to_acceptance == 91 and "median of 20 recent papers" in jms.sources["days_to_acceptance"]
    carbon = res.ranked[0].journal.metrics
    assert carbon.days_to_first_decision == 28 and carbon.acceptance_rate == 25 and carbon.quartile == "Q1"
    assert "insights" in carbon.sources["days_to_first_decision"]
    diag = "\n".join(res.trace)
    assert "Carbon (Elsevier BV)" in diag and "metrics on page: days_to_first_decision=28" in diag
    assert "link.springer.com/journal/99999" in diag


def test_report_roundtrip(tmp_path, monkeypatch):
    from journalfit.report import candidate_from_dict, write_report

    monkeypatch.setattr("journalfit.config.OVERRIDES_PATH", tmp_path / "none.csv")
    p = tmp_path / "paper.docx"
    make_docx(p)
    idx = IndexingData()
    idx.scimago = parse_scimago(SCIMAGO)
    finder = Finder(FakeHttp(), indexing=idx, use_browser=False, log=lambda *_: None)
    res = finder.find(manuscript.parse(p), Criteria(scopus=True), top=3)
    write_report(res, tmp_path / "out")
    data = json.loads((tmp_path / "out" / "shortlist.json").read_text())
    c = candidate_from_dict(data["ranked"][0])
    assert c.journal.name == res.ranked[0].journal.name
    assert "## 1." in (tmp_path / "out" / "shortlist.md").read_text()
    assert "fetched https://www.sciencedirect.com" in (tmp_path / "out" / "diagnostics.txt").read_text()


SCIMAGO_2025 = (
    "Rank;Sourceid;Title;Type;Issn;Publisher;Open Access;Open Access Diamond;SJR;SJR Best Quartile;H index;"
    "Total Docs. (2025);Country;Region;Publisher;Coverage;Categories\n"
    '1;1;Carbon;journal;"00086223, 18733891";Elsevier;No;No;2,171;Q1;300;900;UK;Europe;Elsevier;1963-2025;Materials (Q1)\n'
    '2;2;Batteries;journal;"23130105";MDPI;Yes;No;0,9;Q2;40;300;CH;Europe;MDPI;2015-2025;Energy (Q2)\n'
    '3;3;Old Journal;journal;"11112222";X;No;No;0,3;Q3;20;0;US;NA;X;1990-2001, 2003-2019;Misc (Q3)\n'
)


def test_scimago_2025_layout_oa_and_coverage():
    data = parse_scimago(SCIMAGO_2025)
    assert data["00086223"].quartile == "Q1" and data["00086223"].is_oa is False
    assert data["23130105"].is_oa is True
    assert data["00086223"].active is True and data["11112222"].active is False
    idx = IndexingData()
    idx.scimago = data
    j = Journal("S", "Batteries", issns=["2313-0105"])
    idx.annotate(j)
    assert j.is_oa is True and j.scopus_active is True
    old = Journal("O", "Old Journal", issns=["1111-2222"])
    idx.annotate(old)
    f = Finder.__new__(Finder)
    assert f.apply_static_filters(Candidate(journal=old), Criteria(scopus=True), set()) == "discontinued in Scopus"


def test_ai_lookup_limited_and_cached(tmp_path, monkeypatch):
    from journalfit.finder import MetricCache

    monkeypatch.setattr("journalfit.config.OVERRIDES_PATH", tmp_path / "none.csv")
    monkeypatch.setattr("journalfit.config.DATA_DIR", tmp_path)

    class FakeLLM:
        calls = []

        def lookup_metrics(self, j, pages=None):
            self.calls.append(j.name)
            return Metrics(days_to_first_decision=40.0, sources={"days_to_first_decision": "ai"})

    llm = FakeLLM()
    f = Finder(FakeHttp(), llm=llm, indexing=IndexingData(), use_browser=False, log=lambda *_: None)
    cands = [Candidate(journal=Journal(f"X{i}", f"J{i}")) for i in range(5)]
    f.collect_metrics(cands, Criteria(), web_lookup=True, lookup_top=2)
    assert llm.calls == ["J0", "J1"]  # only the best two
    assert cands[0].journal.metrics.days_to_first_decision == 40 and cands[4].journal.metrics.days_to_first_decision is None
    assert MetricCache().get("X0").days_to_first_decision == 40

    again = [Candidate(journal=Journal("X0", "J0"))]
    f.collect_metrics(again, Criteria(), web_lookup=True, lookup_top=2)
    assert llm.calls == ["J0", "J1"] and again[0].journal.metrics.days_to_first_decision == 40  # from cache


def test_cache_keeps_only_real_findings(tmp_path):
    from journalfit.finder import MetricCache

    c = MetricCache(tmp_path / "c.json")
    c.put("empty", Metrics())
    c.put("full", Metrics(acceptance_rate=20.0, sources={"acceptance_rate": "x"}))
    assert c.get("empty") is None and c.get("full").acceptance_rate == 20


def test_fetcher_rejects_stubs_wrong_pages_and_respects_blocks():
    from journalfit.metrics.publishers import PageFetcher, journal_markers

    class Http:
        pages = {
            "https://ok": _page("Carbon", ""),
            "https://stub": "<html><body>Redirecting...</body></html>",
            "https://other": _page("Some Other Journal", ""),
            "https://botcheck": "<html><title>Just a moment...</title></html>",
        }

        def get_text(self, url, **kw):
            if url == "https://forbidden":
                raise FetchError(f"{url}: HTTP 403")
            return self.pages[url]

    remote_calls = []

    def remote(url):
        remote_calls.append(url)
        return "Carbon. " + "Journal information text. " * 100

    f = PageFetcher(Http(), use_browser=False, remote=remote)
    expect = journal_markers(Journal("S", "Carbon", issns=["0008-6223"]))
    assert f.get("https://ok", expect)
    assert f.get("https://other", expect) is not None  # wrong page locally -> remote reader's copy is used
    assert remote_calls == ["https://other"]
    assert f.get("https://stub", expect) is not None and remote_calls[-1] == "https://stub"
    # Sites that refuse automated access are not fetched another way.
    assert f.get("https://forbidden", expect) is None
    assert f.get("https://botcheck", expect) is None
    assert remote_calls == ["https://other", "https://stub"]


def test_save_override_adds_replaces_and_validates(tmp_path):
    from journalfit.metrics.publishers import load_overrides, save_override

    path = tmp_path / "o.csv"
    save_override("1470-6423", {"days_to_first_decision": "30", "acceptance_rate": ""}, "site", path)
    save_override("0022-2461", {"impact_factor": 4.5}, "", path)
    save_override("14706423", {"acceptance_rate": "15"}, "editor", path)  # same journal, other ISSN format
    data = load_overrides(path)
    assert data["14706423"].acceptance_rate == 15 and data["14706423"].days_to_first_decision is None
    assert data["00222461"].impact_factor == 4.5
    assert "editor" in data["14706423"].sources["acceptance_rate"]
    with pytest.raises(ValueError):
        save_override("1470-6423", {"acceptance_rate": "150"}, "", path)
    with pytest.raises(ValueError):
        save_override("", {"acceptance_rate": "10"}, "", path)


def test_crossref_dates():
    from journalfit.sources.crossref import acceptance_days, parse_date

    import datetime as dt
    assert parse_date("12 March 2024") == dt.date(2024, 3, 12)
    assert parse_date("2024-03-12") == dt.date(2024, 3, 12)
    assert parse_date("March 12th, 2024") == dt.date(2024, 3, 12)
    assert parse_date("12 Mar 2024.") == dt.date(2024, 3, 12)
    assert parse_date("sometime") is None
    items = [
        {"assertion": [{"name": "Received", "value": "1 Jan 2024"}, {"name": "Accepted", "value": "11 Jan 2024"}]},
        {"assertion": [{"name": "received", "value": "1 Jan 2024"}]},  # no accepted date
        {"assertion": [{"name": "received", "value": "1 Jan 2024"}, {"name": "accepted", "value": "1 Jan 2020"}]},  # typo
    ]
    assert acceptance_days(items) == [10]


def test_too_few_dated_papers_gives_nothing():
    from journalfit.sources.crossref import review_speed

    class Http:
        def get_json(self, url, params=None, ttl=None):
            return {"message": {"items": CROSSREF_ITEMS[:5]}}

    assert review_speed("1234-5678", Http()) is None


def test_acceptance_time_rule():
    f = Finder.__new__(Finder)
    crit = Criteria(max_acceptance_days=180)
    assert f.apply_metric_filters(_cand(days_to_acceptance=280.5), crit).startswith("acceptance takes 280.5 days")
    assert f.apply_metric_filters(_cand(days_to_acceptance=145.0), crit) == ""
    unknown = _cand()
    assert f.apply_metric_filters(unknown, crit) == "" and unknown.warnings == ["time to acceptance unknown"]
    assert "accepted within 180 days" in crit.describe()


def test_source_urls_are_tidied():
    from journalfit.llm import _clean_url

    assert _clean_url("https://www.sciencedirect.com/journal/tourism-management?_gl=1%2Aie3zsp%2A_ga%2ANjM") == \
        "https://www.sciencedirect.com/journal/tourism-management"
    assert _clean_url("https://link.springer.com/journal/11356?page=4#1") == "https://link.springer.com/journal/11356"
    assert _clean_url("https://www.tandfonline.com/action/journalInformation?show=journalMetrics&journalCode=x") == \
        "https://www.tandfonline.com/action/journalInformation?show=journalMetrics&journalCode=x"


def test_exact_quartile_choice():
    f = Finder.__new__(Finder)
    j = lambda q: Candidate(journal=Journal("S", "J", metrics=Metrics(quartile=q)))
    only_q2 = Criteria(quartiles=["Q2"])
    assert f.apply_static_filters(j("Q1"), only_q2, set()).startswith("Q1 not among the chosen quartiles")
    assert f.apply_static_filters(j("Q2"), only_q2, set()) == ""
    q1_q3 = Criteria(quartiles=["Q1", "Q3"])
    assert f.apply_static_filters(j("Q3"), q1_q3, set()) == "" and f.apply_static_filters(j("Q2"), q1_q3, set())
    assert f.apply_static_filters(j(None), Criteria(quartiles=["Q2"], if_missing="skip"), set()) == "quartile unknown"
    assert "quartile Q1 or Q3" in q1_q3.describe()


def test_declarations_are_read_from_the_manuscript():
    from journalfit.manuscript import extract_declarations

    lines = ["5. Conclusion", "Tourism matters.", "Declarations",
             "Funding: None.",
             "Conflict of interest",
             "The authors declare that they have no conflict of interest.",
             "Data availability statement",
             "All data are from the World Bank World Development Indicators,",
             "which are publicly available.",
             "Author contributions: I.Y. designed the study and wrote the paper.",
             "Acknowledgements", "We thank the referees."]
    d = extract_declarations(lines)
    assert d["funding"] == "None."
    assert d["competing_interests"] == "The authors declare that they have no conflict of interest."
    assert d["data_availability"].endswith("which are publicly available.")
    assert d["author_contributions"].startswith("I.Y. designed")
    assert "ethics" not in d and d["acknowledgements"] == "We thank the referees."


def test_declarations_written_as_sentences_with_a_full_stop():
    from journalfit.manuscript import extract_declarations

    d = extract_declarations([
        "Declarations",
        "Informed consent. Informed consent was obtained from all participants.",
        "Ethics approval. The study was approved by the college's review board.",
        "Funding. This research received no external funding.",
        "Declaration of generative AI. During the preparation of this manuscript, the authors used Claude for editing.",
        "Data availability is discussed in Section 3. We describe the sample there."])
    assert d["ethics"].startswith("Informed consent was obtained") and "review board" in d["ethics"]
    assert d["funding"] == "This research received no external funding."
    assert d["ai_disclosure"].startswith("During the preparation")
    assert "data_availability" not in d  # a body sentence, not a heading


def test_aims_and_scope_passage():
    from journalfit.metrics.publishers import extract_aims

    page = ("Journal home Aims and scope Tourism Economics publishes applied economic research on tourism, "
            "including demand modelling, tourism and growth, and the environmental effects of tourism, using "
            "rigorous empirical and econometric methods. Editorial board Jane Doe")
    aims = extract_aims(page)
    assert aims.startswith("Tourism Economics publishes") and "Editorial board" not in aims
    assert extract_aims("Nothing relevant here") == ""


def test_backups_and_closest_papers(tmp_path, monkeypatch):
    monkeypatch.setattr("journalfit.config.OVERRIDES_PATH", tmp_path / "none.csv")
    monkeypatch.setattr("journalfit.config.DATA_DIR", tmp_path)

    class LLM:
        prompts = []

        def review_scope(self, ms, cands):
            for c in cands:
                c.scope_fit = 0.8

        def lookup_metrics(self, j, pages=None):
            return Metrics()

    p = tmp_path / "paper.docx"
    make_docx(p)
    idx = IndexingData()
    idx.scimago = parse_scimago(SCIMAGO + '4;4;Slow Journal;journal;"12345679";1;Q1;50;Springer;X\n')
    f = Finder(FakeHttp(), llm=LLM(), indexing=idx, use_browser=False, log=lambda *_: None)
    res = f.find(manuscript.parse(p), Criteria(scopus=True, max_first_decision_days=90), top=5)
    assert [b.journal.name for b in res.backups] == ["Slow Journal"]  # good fit, only too slow
    assert res.backups[0].excluded_reason.startswith("first decision 200")
    carbon = next(c for c in res.ranked if c.journal.name == "Carbon")
    assert carbon.example_works and {"title", "year", "doi", "similarity"} <= set(carbon.example_works[0])


def test_ai_metrics_are_kept_only_when_written_on_the_page(monkeypatch):
    from journalfit import llm as llm_mod
    from journalfit.models import Journal

    page = ("Journal metrics. Submission to first decision (median): 6 weeks. "
            "Acceptance rate 18%. Impact Factor 2.4. Volume 4, issue 10.")
    assert llm_mod.stated_in("days_to_first_decision", 42, page)  # 6 weeks
    assert llm_mod.stated_in("acceptance_rate", 18, page)
    assert llm_mod.stated_in("impact_factor", 2.4, page)
    assert not llm_mod.stated_in("days_to_first_decision", 4, page)  # "Volume 4" is not a decision time
    assert not llm_mod.stated_in("acceptance_rate", 10, page)  # "issue 10" is not a rate
    assert not llm_mod.stated_in("citescore", 2.8, page)

    ds = llm_mod.DeepSeek("deepseek-test", "key")
    monkeypatch.setattr(ds, "web_pages", lambda *a, **k: {"https://publisher.example/j": page})
    monkeypatch.setattr(ds, "structured", lambda *a, **k: {
        "days_to_first_decision": 4, "days_to_acceptance": None, "acceptance_rate": 18,
        "impact_factor": 2.4, "citescore": 2.8, "source_url": "https://publisher.example/j"})
    m = ds.lookup_metrics(Journal(openalex_id="S1", name="Test Journal", publisher="Wiley"))
    assert m.days_to_first_decision is None and m.citescore is None  # invented: dropped
    assert m.acceptance_rate == 18 and m.impact_factor == 2.4


def test_remembered_metrics_from_before_the_check_are_looked_up_again(tmp_path):
    import json
    import time

    from journalfit.finder import MetricCache
    from journalfit.models import Metrics

    path = tmp_path / "metric_cache.json"
    path.write_text(json.dumps({"S1": {"ts": time.time(), "metrics": {"days_to_first_decision": 4.0,
                                                                      "sources": {"days_to_first_decision": "x"}}}}))
    cache = MetricCache(path)
    assert cache.get("S1") is None
    cache.put("S2", Metrics(impact_factor=2.0, sources={"impact_factor": "y"}))
    assert cache.get("S2").impact_factor == 2.0
